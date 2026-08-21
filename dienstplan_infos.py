import os
import sys
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import requests
import openpyxl
from openpyxl.styles import Alignment, Font

sys.stdout.reconfigure(encoding="utf-8")

# Manuell vor jedem Start setzen:
# False -> nur Simulation, es wird ausschließlich die Excel-Datei mit den simulierten
#          Werten erzeugt, es finden KEINE Änderungen in ChurchTools statt.
# True  -> die simulierten Änderungen werden live in ChurchTools geschrieben und
#          zusätzlich als Änderungsprotokoll in eine Excel-Datei geschrieben.
APPLY_EDITS = False

BERLIN = ZoneInfo("Europe/Berlin")

URL_NAME = "volksdorf"
BASE_URL = f"https://{URL_NAME}.church.tools/api"
TOKEN = os.getenv("CHURCHTOOLS_TOKEN_VOL")

FROM = date.today().strftime("%Y-%m-%d")
TO = (date.today() + timedelta(days=365)).strftime("%Y-%m-%d")

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "Dienstplan_Infos.xlsx")
CHANGES_OUTPUT_PATH = os.path.join(SCRIPT_DIR, "Dienstplan_Aenderungen.xlsx")

COLUMNS = [
    "Datum", "Uhrzeit", "Kalender", "Titel", "Weitere Infos", "Ressourcenbuchungen",
    "Titel (simuliert)", "Weitere Infos (simuliert)",
]

# Simulierte Umbenennung: alter Titel -> neuer Titel
TITEL_SIMULATION = {
    "Gottesdienst - Rockenhof": "Gottesdienst",
    "Gottesdienst - St. Gabriel": "Gottesdienst",
}

# Simulierte "Weitere Infos": Ressourcenbuchung -> Text, nur wenn "Weitere Infos" leer ist
WEITERE_INFOS_SIMULATION = {
    "Kirche Rockenhof": "am Rockenhof",
    "Kirche St. Gabriel": "in St. Gabriel",
}

README_RULES = [
    "Regel 1: Wenn der Titel \"Gottesdienst - Rockenhof\" ist, wird der Titel auf \"Gottesdienst\" geändert.",
    "Regel 2: Wenn der Titel \"Gottesdienst - St. Gabriel\" ist, wird der Titel auf \"Gottesdienst\" geändert.",
    "Regel 3: Wenn eine Ressourcenbuchung \"Kirche Rockenhof\" vorliegt UND \"Weitere Infos\" leer ist, "
    "wird \"Weitere Infos\" auf \"am Rockenhof\" gesetzt.",
    "Regel 4: Wenn eine Ressourcenbuchung \"Kirche St. Gabriel\" vorliegt UND \"Weitere Infos\" leer ist, "
    "wird \"Weitere Infos\" auf \"in St. Gabriel\" gesetzt.",
]

CHANGES_COLUMNS = [
    "Datum", "Uhrzeit", "Kalender", "Titel (vorher)", "Titel (nachher)",
    "Weitere Infos (vorher)", "Weitere Infos (nachher)", "Status", "Fehler",
]


def _get_all_pages(url, headers, params):
    response = requests.get(url, headers=headers, params=params)
    response.raise_for_status()
    data = response.json().get("data", [])
    pagination = response.json().get("meta", {}).get("pagination")
    if pagination:
        for page in range(2, pagination.get("lastPage") + 1):
            response = requests.get(f"{url}?page={page}", headers=headers, params=params)
            response.raise_for_status()
            data.extend(response.json().get("data", []))
    return data


def get_calendar_ids(headers):
    calendars = _get_all_pages(f"{BASE_URL}/calendars", headers, {})
    return [c.get("id") for c in calendars if c.get("id") is not None]


def get_resource_ids(headers):
    resources = _get_all_pages(f"{BASE_URL}/resources", headers, {})
    return [r.get("id") for r in resources if r.get("id") is not None]


def get_events(headers, calendar_ids, from_date, to_date):
    params = {
        "calendar_ids[]": calendar_ids,
        "from": from_date,
        "to": to_date,
        "include": "eventServices",
        "canceled": "false",
    }
    return _get_all_pages(f"{BASE_URL}/events", headers, params)


def get_bookings(headers, resource_ids, from_date, to_date):
    params = {
        "resource_ids[]": resource_ids,
        "from": from_date,
        "to": to_date,
    }
    return _get_all_pages(f"{BASE_URL}/bookings", headers, params)


def build_entries(events, bookings):
    entries = []
    for event in events:
        start_date = datetime.fromisoformat(event.get("startDate")).astimezone(BERLIN)
        kalender = event.get("calendar", {}).get("title") or "(kein Kalendername)"
        titel = event.get("name") or "(kein Titel)"
        raw_note = event.get("note") or ""
        weitere_infos = raw_note or "(keine weiteren Infos)"

        appointment_id = event.get("appointmentId")
        resource_bookings = [
            b.get("base", {}).get("resource", {}).get("name")
            for b in bookings
            if appointment_id and b.get("base", {}).get("appointmentId") == appointment_id
        ]
        resource_bookings = list(dict.fromkeys(r for r in resource_bookings if r))
        ressourcen = " | ".join(resource_bookings) if resource_bookings else "(keine Ressourcenbuchung)"

        titel_simuliert = TITEL_SIMULATION.get(titel, titel)

        weitere_infos_simuliert = weitere_infos
        if not raw_note:
            for resource_name, simulierter_text in WEITERE_INFOS_SIMULATION.items():
                if resource_name in resource_bookings:
                    weitere_infos_simuliert = simulierter_text
                    break

        calendar_id_raw = event.get("calendar", {}).get("domainIdentifier")

        entries.append({
            "event_id": event.get("id"),
            "appointment_id": appointment_id,
            "calendar_id": int(calendar_id_raw) if calendar_id_raw is not None else None,
            "start_date": start_date,
            "kalender": kalender,
            "titel": titel,
            "weitere_infos": weitere_infos,
            "ressourcen": ressourcen,
            "titel_simuliert": titel_simuliert,
            "weitere_infos_simuliert": weitere_infos_simuliert,
        })
    return entries


def has_diff(entry):
    return entry["titel_simuliert"] != entry["titel"] or entry["weitere_infos_simuliert"] != entry["weitere_infos"]


def entries_to_rows(entries):
    return [
        [
            e["start_date"].strftime("%d.%m.%Y"),
            e["start_date"].strftime("%H:%M"),
            e["kalender"],
            e["titel"],
            e["weitere_infos"],
            e["ressourcen"],
            e["titel_simuliert"],
            e["weitere_infos_simuliert"],
        ]
        for e in entries
    ]


def create_excel(rows, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dienstplan"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col, width in zip("ABCDEFGH", [12, 9, 22, 30, 40, 30, 30, 40]):
        ws.column_dimensions[col].width = width
    wb.save(output_path)


def _build_readme_sheet(wb):
    ws = wb.create_sheet(title="README")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 110

    bold_h1 = Font(bold=True, size=14)
    bold_h2 = Font(bold=True, size=11)
    body = Font(size=11)
    wrap = Alignment(wrap_text=True, vertical="top")

    row = [1]

    def write(value, font=body, height=18):
        c = ws.cell(row=row[0], column=1, value=value)
        c.font = font
        c.alignment = wrap
        ws.row_dimensions[row[0]].height = height
        row[0] += 1

    def blank():
        ws.row_dimensions[row[0]].height = 8
        row[0] += 1

    write("Regeln für automatische Änderungen im ChurchTools-Dienstplan (Volksdorf)", font=bold_h1, height=24)
    blank()
    write("Dieses Blatt beschreibt die Regeln, nach denen dienstplan_infos.py Titel und "
          "\"Weitere Infos\" der Dienstplan-Einträge automatisch anpasst, wenn APPLY_EDITS = True "
          "gesetzt ist.", height=36)
    blank()
    write("Regeln", font=bold_h2, height=18)
    for rule in README_RULES:
        write(rule, height=32)
    blank()
    write("Wichtig: Die Regeln 3 und 4 greifen ausschließlich, wenn \"Weitere Infos\" zum Zeitpunkt "
          "des Skriptlaufs leer ist. Ein bereits vorhandener Wert wird nie überschrieben.", height=36)
    blank()
    write("Technische Umsetzung", font=bold_h2, height=18)
    write("- Titeländerung: PUT /api/calendars/{calendarId}/appointments/{appointmentId}. "
          "Der Titel hängt am Kalendertermin, nicht am Dienstplan-Event, daher wird der komplette "
          "Termin gelesen (Read), nur das Feld 'title' geändert und das volle Objekt zurückgeschrieben "
          "(Write) - inkl. 'version' für ChurchTools' Optimistic Locking.", height=64)
    write("- \"Weitere Infos\"-Änderung: PUT /api/events/{eventId} mit dem Feld 'note' "
          "(zusammen mit den unveränderten aktuellen Werten für 'isCanceled' und 'adminIds').", height=32)
    blank()
    write(f"Ausgeführt am: {datetime.now(BERLIN).strftime('%d.%m.%Y %H:%M')} für die ChurchTools-Instanz "
          "Volksdorf (VOL).", height=18)
    return ws


def create_changes_excel(changes, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Änderungen"
    ws.append(CHANGES_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for c in changes:
        ws.append([
            c["start_date"].strftime("%d.%m.%Y"),
            c["start_date"].strftime("%H:%M"),
            c["kalender"],
            c["titel_vorher"],
            c["titel_nachher"],
            c["infos_vorher"],
            c["infos_nachher"],
            c["status"],
            c["fehler"],
        ])
    for col, width in zip("ABCDEFGHI", [12, 9, 22, 28, 28, 30, 30, 12, 40]):
        ws.column_dimensions[col].width = width

    _build_readme_sheet(wb)
    wb.move_sheet("README", offset=-(len(wb.sheetnames) - 1))
    wb.active = wb["README"]

    wb.save(output_path)


def update_event_note(headers, event_id, new_note):
    response = requests.get(f"{BASE_URL}/events/{event_id}", headers=headers)
    response.raise_for_status()
    current = response.json()["data"]
    body = {
        "note": new_note,
        "isCanceled": current.get("isCanceled", False),
        "adminIds": current.get("adminIds", []),
    }
    put_response = requests.put(f"{BASE_URL}/events/{event_id}", headers=headers, json=body)
    put_response.raise_for_status()


def update_appointment_title(headers, calendar_id, appointment_id, new_title):
    response = requests.get(f"{BASE_URL}/calendars/{calendar_id}/appointments/{appointment_id}", headers=headers)
    response.raise_for_status()
    appointment = response.json()["data"]["appointment"]
    # Nur die aktuellen (nicht-deprecated) Felder zurückschreiben
    for key in ("@deprecated", "caption", "note", "information", "additions"):
        appointment.pop(key, None)
    appointment["title"] = new_title
    put_response = requests.put(
        f"{BASE_URL}/calendars/{calendar_id}/appointments/{appointment_id}",
        headers=headers,
        json=appointment,
    )
    put_response.raise_for_status()


def apply_changes(headers, entries):
    changes = []
    for e in entries:
        if not has_diff(e):
            continue
        needs_title_change = e["titel_simuliert"] != e["titel"]
        needs_infos_change = e["weitere_infos_simuliert"] != e["weitere_infos"]

        change = {
            "start_date": e["start_date"],
            "kalender": e["kalender"],
            "titel_vorher": e["titel"],
            "titel_nachher": e["titel_simuliert"] if needs_title_change else e["titel"],
            "infos_vorher": e["weitere_infos"],
            "infos_nachher": e["weitere_infos_simuliert"] if needs_infos_change else e["weitere_infos"],
            "status": "OK",
            "fehler": "",
        }
        try:
            if needs_title_change:
                update_appointment_title(headers, e["calendar_id"], e["appointment_id"], e["titel_simuliert"])
            if needs_infos_change:
                update_event_note(headers, e["event_id"], e["weitere_infos_simuliert"])
        except Exception as exc:
            change["status"] = "Fehler"
            change["fehler"] = str(exc)
        changes.append(change)
    return changes


def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--from", dest="from_date", default=FROM, help="Start date YYYY-MM-DD")
    parser.add_argument("--to", dest="to_date", default=TO, help="End date YYYY-MM-DD")
    args = parser.parse_args()

    if not TOKEN:
        print("Das Access Token ist nicht gesetzt. Bitte die Umgebungsvariable 'CHURCHTOOLS_TOKEN_VOL' definieren.")
        sys.exit(1)

    headers = {"Authorization": f"Login {TOKEN}"}

    calendar_ids = get_calendar_ids(headers)
    print(f"{len(calendar_ids)} Kalender gefunden: {calendar_ids}")

    resource_ids = get_resource_ids(headers)
    print(f"{len(resource_ids)} Ressourcen gefunden.")

    events = get_events(headers, calendar_ids, args.from_date, args.to_date)
    print(f"{len(events)} Dienstplan-Einträge gefunden.")

    bookings = get_bookings(headers, resource_ids, args.from_date, args.to_date)
    print(f"{len(bookings)} Ressourcenbuchungen gefunden.")

    entries = build_entries(events, bookings)
    changed_entries = [e for e in entries if has_diff(e)]
    rows = entries_to_rows(changed_entries)

    create_excel(rows, OUTPUT_PATH)
    print(f"{len(changed_entries)} Einträge mit Abweichungen (von {len(entries)} insgesamt).")
    print(f"Excel-Datei gespeichert: {OUTPUT_PATH}")

    if APPLY_EDITS:
        print("\nAPPLY_EDITS = True -> Änderungen werden jetzt live in ChurchTools durchgeführt ...")
        changes = apply_changes(headers, entries)
        n_ok = sum(1 for c in changes if c["status"] == "OK")
        n_err = len(changes) - n_ok
        print(f"{len(changes)} Änderungen durchgeführt ({n_ok} erfolgreich, {n_err} fehlgeschlagen).")
        create_changes_excel(changes, CHANGES_OUTPUT_PATH)
        print(f"Änderungsprotokoll gespeichert: {CHANGES_OUTPUT_PATH}")
    else:
        print("\nAPPLY_EDITS = False -> es wurden keine Änderungen in ChurchTools durchgeführt (nur Simulation in Excel).")


if __name__ == "__main__":
    main()
